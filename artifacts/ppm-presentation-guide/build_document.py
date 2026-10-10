from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'Enterprise PPM Presentation Guide.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.72)
sec.footer_distance = Inches(.28)
for style in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    s = doc.styles[style]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.widow_control = True
    for el in s.element.xpath('./w:pPr/w:pBdr'):
        el.getparent().remove(el)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Subtitle'].font.size = Pt(13)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].font.size = Pt(21)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(11)
doc.styles['Heading 2'].font.size = Pt(12.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
for style in ['List Bullet', 'List Number']:
    doc.styles[style].font.size = Pt(11)
    doc.styles[style].paragraph_format.space_after = Pt(4)
    doc.styles[style].paragraph_format.line_spacing = 1.08
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Enterprise PPM Presentation Guide  |  ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Enterprise PPM Presentation Guide'
doc.core_properties.subject = 'Section functions required information and long term value'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Enterprise PPM, presentation, project management, user guide'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    a = doc.add_heading(title, 1)
    a.paragraph_format.page_break_before = True

def lead(label, text):
    a = p('')
    a.add_run(label + ' ').bold = True
    a.add_run(text)
    return a

def bullets(items):
    for item in items:
        p(item, 'List Bullet')

def numbered(items):
    for i, item in enumerate(items, 1):
        a = p('')
        a.add_run(f'{i}. ').bold = True
        a.add_run(item)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    pr.append(borders)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = text
    for ri, row in enumerate(t.rows):
        nosplit = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(nosplit)
        for ci, c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = c._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '23374D' if ri == 0 else ('F1F4F7' if ri % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '75'), ('bottom', '75'), ('left', '100'), ('right', '100')]:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), value); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            for a in c.paragraphs:
                a.paragraph_format.space_after = Pt(0)
                a.paragraph_format.line_spacing = 1.05
                for r in a.runs:
                    r.font.size = Pt(10)
                    if ri == 0:
                        r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    p('').paragraph_format.space_after = Pt(0)
    return t

def explain(name, function, info, reason):
    h(name)
    p(function)
    lead('Information to enter', info)
    lead('Why it is needed', reason)

doc.add_paragraph('Enterprise PPM Presentation Guide', 'Title')
doc.add_paragraph('Section functions required information and long term value', 'Subtitle')
p('10 October 2026')
p('Enterprise PPM helps an organization plan projects, manage people and money, record risks, and review whether its work supports business goals. PPM means Project and Portfolio Management. This guide explains the workspace in plain language for presenters and audiences who are new to project management.')
p('Each section explains its purpose, the information entered, why that information matters, and the value it can provide over time. The presentation sentences can be read aloud or adapted into slides. Long term benefits depend on people keeping the records accurate and current.')
h('The nine menu sections')
table(['Section', 'Main question it helps answer'], [
    ['Workspace overview', 'Where should I start and what needs my attention?'],
    ['Projects', 'What are we delivering and how is the work progressing?'],
    ['Maturity assessments', 'How consistently do we follow useful working practices?'],
    ['Command center', 'Where do managers need to investigate or act?'],
    ['Integrations and jobs', 'How do we bring in supported data and check the outcome?'],
    ['Resource planning', 'Do people have enough time for their commitments?'],
    ['Portfolios', 'How does our group of projects support business priorities?'],
    ['Account and access', 'Who can see or change which information?'],
    ['Component gallery', 'How do we keep the application screens consistent?'],
], [2.05, 4.78])
h('How to read the field guidance')
p('Required means the record or action needs that information. Optional information can still be important for reporting. Most business records require a Name and offer a Description. Selecting an Owner identifies responsibility; permissions must still be granted separately.')
p('The implemented features use saved records, defined calculations, and human review. AI analysis, predictive models, and general Word or PDF document intake are planned capabilities.')

page('1 Workspace overview')
lead('Presentation sentence', 'This is the home page. It helps each person understand where to begin and which responsibilities need attention.')
p('The workspace combines navigation, setup progress, and personal work. A new user can start here without needing to understand every module at once.')
h('What is inside')
bullets([
    'Getting started shows saved setup progress and links to the next steps.',
    'My work links to permitted assigned work, open risks, and active assessments.',
    'Membership information shows the role in the selected organization.',
    'Connected organizations and active sessions help the user understand their account context.',
    'Quick links lead to projects, profile and security controls, and member management where permitted.',
    'In-app notifications show eligible assignment, review, and import notices. Users can mark them read and control their preference. Eligible notices are displayed for 30 days.',
])
h('Information needed and why')
p('This page mainly reads information saved elsewhere. The signed-in account, selected organization, role, access grants, assignments, risk ownership, and saved setup records determine what appears. There is no separate project plan to fill in on the home page.')
lead('Example', 'Priya signs in and sees an assigned assessment and an open risk she owns. She can open the relevant records directly.')
lead('Long term value', 'A consistent starting point makes onboarding easier and helps people find their responsibilities as the number of projects grows. Its usefulness depends on assignments and ownership being maintained.')
h('Related organization setup')
p('Organization setup is linked from the application even though it is not a menu item in the supplied screenshot. An Admin uses it to establish the common settings that planning depends on.')
table(['Setup information', 'Why it is needed'], [
    ['Timezone and three-letter currency code', 'Establish the local date context and financial reporting currency. These settings are needed before delivery records can be created.'],
    ['Departments teams and team membership', 'Describe who does the work and how people are grouped.'],
    ['Strategic goals', 'Record the business priorities that projects and portfolios can support.'],
    ['Project and portfolio grants', 'Give accepted members the record access needed for their responsibilities.'],
    ['Metric policy', 'Sets validated risk thresholds and source freshness rules used by dashboards.'],
], [2.2, 4.63])

page('2 Projects and project information')
lead('Presentation sentence', 'This section records what we are delivering, who is responsible, the dates and money involved, and anything that could affect progress.')
p('The Projects list helps users find permitted projects, search and filter records, and open a detailed project workspace. Inside a project, tabs separate delivery work, financial records, risks, benefits, and reference plans.')
h('Information in the project form')
table(['Information', 'Requirement and purpose'], [
    ['Name and description', 'Name is required. Use the description to explain the objective and intended result.'],
    ['Portfolio and program', 'Portfolio is required. Program is optional and links related projects pursuing a shared result.'],
    ['Department team and owner', 'Identify the delivery group and accountable person. These selections do not automatically grant access.'],
    ['Status', 'Choose draft, active, on hold, completed, or cancelled to describe the current stage.'],
    ['Reported health', 'Choose Unknown, Green, Amber, or Red to record the manager\'s assessment. Explain the reason in status updates.'],
    ['Currency', 'A three-letter currency code is required so amounts have an unambiguous meaning.'],
    ['Start and planned finish', 'Record the intended schedule. Not every date is mandatory at initial creation, but schedule measures need appropriate dates.'],
    ['Forecast finish', 'Records the current estimate of when the project will finish.'],
    ['Actual start and finish', 'Records what really happened. Enter actual dates only when the event has occurred.'],
    ['Expenditure ledger complete', 'Indicates whether spending records are complete enough for supported financial measures.'],
    ['Optional classification', 'Priority, industry, category, working method, and size help describe and filter projects.'],
], [2.0, 4.83])
h('Three kinds of dates')
p('Planned is the agreed target. Forecast is the current estimate. Actual is the event that has already happened. Keeping these separate makes a delay visible without erasing the original commitment.')
lead('Example', 'A customer portal is planned to finish on 6 November. Its forecast moves to 13 November. The project record shows the revised expectation, while a saved baseline retains the reference plan.')
lead('Long term value', 'A shared project record supports handovers and gives future teams useful evidence about commitments, changes, spending, and the causes of delays.')

page('Projects delivery work and checkpoints')
p('These tabs turn a broad objective into manageable work. Most records require a Name and a Project link, and offer a Description. The project can already be selected when creating the record inside its workspace.')
explain('Work items',
    'A work item is a small piece of delivery work. The work board groups records into backlog, ready, in progress, and done.',
    'Project and name; status; optional owner, sprint, start date, due date, priority, work type, blocked flag, and story points.',
    'These fields explain what needs doing, who should do it, and its progress. Story points are a team-specific estimate, not a number of working hours.')
explain('Milestones',
    'A milestone is an important checkpoint, such as requirements agreed, testing finished, or launch ready. Milestones also support the delivery timeline.',
    'Project, name, and due date are required. Add actual finish when achieved and mark whether the milestone is critical.',
    'Managers need to see important commitments and distinguish a missed critical checkpoint from an ordinary task delay.')
explain('Sprints',
    'A sprint is a short period in which a team plans and reviews a group of work items.',
    'Project, team, start date, end date, and name are required. Choose planned, active, or closed status.',
    'The period and team give the work a clear context. Closing a sprint preserves its recorded completion information; the closed sprint cannot simply be reopened.')
explain('Blockers',
    'A blocker is an obstacle that is already preventing progress.',
    'Project and name; description of the obstacle; optional related work item and owner; open or resolved status.',
    'Recording the obstacle and responsible person helps the team discuss what must change before work can continue.')
explain('Dependencies',
    'A dependency shows that one project or activity relies on another. A predecessor comes first; a successor follows it.',
    'Source project, successor project, and name are required. Link relevant predecessor and successor work items or milestones where applicable.',
    'Teams need to coordinate handoffs. If testing depends on another project\'s API release, that relationship should be visible before a delay spreads.')
lead('Long term value', 'These records reveal which tasks and handoffs repeatedly cause delays. They also let a replacement manager understand the work without relying entirely on the previous manager\'s memory.')

page('Projects risks money and reference plans')
explain('Risk register',
    'A risk is something that might cause a future problem. A blocker is already stopping work.',
    'Name, project, severity from 1 to 5, likelihood from 1 to 5, and review date. Add owner, mitigation, status, and resolution or acceptance reason.',
    'Severity records impact; likelihood records how likely the problem is. Mitigation means the action to reduce it. Moving to mitigating needs an owner and plan; acceptance or closure needs an owner and explanation.')
explain('Budget lines and expenditure',
    'Budget lines record funding. Expenditure records money actually spent.',
    'Each needs name, project, amount, and currency. Budget lines include an approved flag. Expenditure also requires spending date; a reversal references the original entry.',
    'Separating planned funding from actual spending makes the financial position understandable. Expenditure history is preserved; corrections use supported reversal entries rather than rewriting the original.')
explain('Business benefits',
    'Benefits describe the useful result expected from the project, such as hours saved or money saved.',
    'Name, project, unit, target, and period start/end are required. Enter realized value when known; blank means unknown.',
    'The unit and period explain what is being measured. This lets the organization review value after delivery instead of considering completion alone.')
explain('Status updates and plan baselines',
    'A status update preserves a progress explanation. A baseline preserves a fixed reference copy of the plan.',
    'Status update: project, name, description, and reported health. Baseline: project, name, and required baseline reason.',
    'New updates preserve the reporting history. A baseline allows later changes to be compared with a reference plan rather than making the original commitment disappear.')
explain('Metric plans goal alignment and access',
    'These tabs supply calculation inputs, connect the project to business priorities, and control record permissions.',
    'Metric plan requires project, baseline, and working calendar; forecast final cost is optional. Alignment needs one project or portfolio scope, a goal, weight, and contribution. Access needs a member, one scope, and read/edit/manage level.',
    'Calculations need explicit reference inputs. Alignment weights must total 1, meaning 100%. Admin-managed grants allow people to work with the intended records.')
lead('Long term value', 'Together these records preserve the reasons for decisions, provide a traceable financial record, and support learning about which projects deliver useful results.')

page('3 Maturity assessments and framework design')
lead('Presentation sentence', 'This section checks whether a team follows useful working practices consistently. It combines clear questions, supporting evidence, and independent review.')
p('Maturity describes how consistently a team works. A team can deliver one successful project while still having weak planning or risk practices. An assessment checks the practices behind the result.')
h('The framework library')
p('A framework is the agreed set of questions and scoring rules. Admin creates and publishes a framework. A Project Manager with an explicit framework-editor grant can edit its draft. Publishing creates a version that existing assessments continue to use.')
table(['Framework information', 'Why it is needed'], [
    ['Name description and category', 'Identify the purpose and subject. Categories include custom, PMO, agile, governance, and delivery. PMO refers to the project management office.'],
    ['Dimension name and weight', 'A dimension is a group of questions, such as planning. Its weight states its importance in the overall score. Dimension weights must total 1, or 100%.'],
    ['Question name and prompt', 'State what the respondent must evaluate. Clear wording reduces different interpretations.'],
    ['Answer type', 'Choose yes/no, rubric level, or text according to the information required.'],
    ['Question weight', 'A positive weight controls the question\'s relative importance within its dimension.'],
    ['Required flag', 'Identifies questions that must be answered before submission and completion.'],
    ['Rubric values and labels', 'Define numeric levels with agreed meanings. A rubric needs at least two distinct levels.'],
], [2.08, 4.75])
h('Answer types in easy words')
bullets([
    'Yes/no asks whether a practice exists, such as whether a project has a named owner.',
    'Rubric asks how well a practice is demonstrated. Example levels might be 0 Not demonstrated, 1 Partly demonstrated, and 2 Established and repeatable.',
    'Text collects a written explanation. Text answers contribute to completion and answer coverage, but are not numerically scored.',
])
p('Each dimension needs at least one scored question before publication. Evidence explains an answer; attaching evidence does not automatically add score points or prove the claim true.')
lead('Long term value', 'An agreed framework makes reviews more consistent. Preserved versions make it possible to explain the rules behind old results even after the framework changes.')

page('Maturity assignments evidence and review')
h('Assign an assessment')
p('Admin selects a published framework version and creates an assignment. Assessment name, published version, respondent, and independent reviewer are needed. Department and team can identify the assessed scope. Respondent and reviewer must be different active Admin or Project Manager members.')
p('Respondent means the person who answers. Independent reviewer means the different person who checks those answers. These are assessment assignments, not additional organization roles.')
h('Answer and provide evidence')
p('The respondent saves the required answers and supporting evidence notes. Evidence can describe dated records, decisions, or review practices. Supported attachments are currently bounded plain-text evidence with a release process. General Word/PDF intake and AI document analysis are planned.')
lead('Example question', 'Are project risks recorded and reviewed regularly? A useful evidence note identifies the risk register, review dates, and the person responsible. It should support the answer actually selected.')
h('Submit and independently review')
p('The respondent submits the completed assessment. The assigned reviewer checks answers and supporting evidence, adds a review note, and completes the review. If an allowed numeric adjustment is made, a reason is required and the original answer is retained. The reviewer can use supported return-for-rework actions when corrections are needed.')
h('Read results and coverage')
table(['Result', 'Easy explanation'], [
    ['Dimension and overall scores', 'Defined calculations summarize the scored answers and valid reviewer adjustments. They use a 0 to 100 scale.'],
    ['Answer coverage', 'Shows how much of the question weight has valid answers. A provisional score can exist even while required answers are missing.'],
    ['Evidence coverage', 'Shows how many questions have qualifying evidence. It measures presence, not whether the evidence is convincing.'],
    ['History', 'Preserves recorded assessment actions and review decisions.'],
    ['Department comparison', 'Compares permitted completed results for a selected framework version and period. Comparable scope and rules matter.'],
], [1.85, 4.98])
h('Share results with the right people')
p('A completed-result grant can allow a Viewer to inspect scores. It does not share private answers or evidence. Viewers cannot submit assessments, evidence, or reviews. Framework editing, assessment participation, and result access are separate permissions.')
lead('Long term value', 'Repeated assessments can show whether practices improve. The team can identify a weak area, agree an improvement, and check later whether evidence supports better practice. A score is a review aid, not a guarantee of project success.')

page('4 Command center')
lead('Presentation sentence', 'This dashboard brings project, money, risk, and staffing information together so managers can see where attention is needed.')
p('The dashboard mainly reads saved information from other modules. Users choose a period, permitted portfolio, and dashboard view. They can save personal filter selections under a view name. Permitted editors can also record risks or benefits through dashboard actions.')
table(['View', 'Information shown and purpose'], [
    ['Executive', 'Project health, financial measures, goal alignment, and critical milestones support an overall management review.'],
    ['Delivery', 'Schedule differences, milestone measures, unfinished overdue critical checkpoints, and closed-sprint velocity help investigate progress.'],
    ['Risk', 'Schedule, financial, and capacity risk components, plus recorded risks and delivery issues, help explain the sources of concern.'],
    ['Resource', 'Assigned-resource capacity, allocations, excess hours, and permitted team signals help identify staffing pressure.'],
], [1.2, 5.63])
h('Information needed for the measures')
bullets([
    'Schedule measures need appropriate project dates, a frozen baseline, and an explicit working calendar.',
    'Spending measures need approved budget and expenditure records. Budget usage needs a complete ledger and a nonzero budget.',
    'Financial risk uses a forecast final cost as well as the approved budget.',
    'Capacity measures need availability, calendars, time off, allocations, and sufficient permitted scope.',
    'Goal alignment uses entered goals, weights, and contributions.',
])
h('How to explain health and Unknown')
p('Reported health is the manager\'s selected status. Evidence health is calculated using defined rules. The current combined risk uses schedule, financial, and capacity components with weights of 40%, 35%, and 25%. All three must be available for that combined result.')
p('Unknown means there is not enough information or permitted scope for a reliable calculation. It must not be read as zero risk, zero spending, or a healthy project. Risk scores are rule-based signals, not probabilities of failure.')
lead('Example', 'A project is reported Green, but an unfinished critical milestone is overdue. The dashboard can flag the disagreement so the manager can inspect the source records.')
h('Source explanations and reporting dates')
p('Measures include their period, source facts, coverage, freshness, and explanation. Reloading a dashboard does not make old source facts fresh. The selected dates filter relevant period information; they do not reconstruct every detail of the project as it existed in the past. Money remains separated by currency.')
lead('Long term value', 'Managers can review projects consistently and investigate the facts behind a warning. Visible coverage and source age also encourage better reporting habits.')

page('5 Integrations and jobs')
lead('Presentation sentence', 'This section brings supported information from other tools into the workspace and shows whether the import worked.')
p('An integration is a configured connection. A job is a processing run, such as importing a batch of work items. Admin manages connections and starts the supported import actions.')
h('Information used to create a connection')
table(['Information', 'Purpose and condition'], [
    ['Connection name and project', 'Required to identify the connection and destination for the imported records.'],
    ['Provider and mode', 'Select CSV, GitHub, or Jira. Fixture mode uses supplied test records; live mode calls a provider where supported and configured.'],
    ['Repository or Jira Cloud host', 'Identifies the source for the applicable provider. Jira also uses a project key where required.'],
    ['Entity type', 'Chooses issues/work items, GitHub pull requests, or Jira sprints where supported.'],
    ['Field mapping', 'Required mapping explains which source field provides each local value. For example, a source title becomes a work-item name.'],
    ['Status mapping', 'Translates source labels into local statuses. For example, To Do can map to backlog.'],
    ['Jira board and local team', 'Needed for the applicable sprint import configuration.'],
    ['Server secret reference', 'Points to privately configured credentials for live access. It is a reference, not the password or token itself.'],
    ['Webhook secret reference', 'Supports signed incoming updates for the applicable connection.'],
    ['CSV or fixture records', 'Supplies supported input for preview and fixture/import processing.'],
], [2.12, 4.71])
h('What users do inside the section')
numbered([
    'Create the connection and check its destination and mappings.',
    'Preview supported supplied data to identify problems before importing.',
    'Start the job and follow its processing stages.',
    'Inspect outcomes, rejected records, and available data-quality explanations.',
    'Review sync history and use supported connection or cancellation controls when needed.',
])
p('Cancelling a job does not undo effects already saved. Repeated imports within the same connection use source identities to avoid simply creating another copy; separate duplicate connections still need careful configuration.')
lead('Long term value', 'Supported imports can reduce repeated manual entry and preserve the origin of records. Visible progress and rejection details make failures easier to investigate.')
lead('Current capability', 'CSV imports and labelled Jira/GitHub fixtures are implemented. Actual live Jira/GitHub acceptance remains unverified without provider credentials and acceptance checks. Successful fixtures do not establish a working customer connection.')

page('6 Resource planning foundations')
lead('Presentation sentence', 'This section checks whether people have enough working time for their commitments, including holidays, leave, and work on other projects.')
p('A resource profile describes a person for staffing. It can exist without a login account. Creating a resource does not invite that person to the website or automatically grant account access.')
h('Set up the people and their working time')
table(['Record', 'Information to enter', 'Why it is needed'], [
    ['Resource profile', 'Required Name; description; optional linked login member.', 'Identifies the person whose staffing is planned.'],
    ['Working calendar', 'Name, timezone, hours for all seven weekdays, and holiday dates where applicable.', 'Defines normal scheduled hours and nonworking holidays.'],
    ['Availability', 'Name, resource, start date, excluded end date, and calendar version.', 'States which schedule applies to the person during the period.'],
    ['Skills', 'Name, resource, and skill.', 'Helps find people with relevant capabilities.'],
    ['Resource teams', 'Name, resource, team, start date, and excluded end date.', 'Shows dated team membership for staffing.'],
    ['Approved time off', 'Name, resource, dates, and hours off per scheduled day.', 'Reduces available time for approved absence.'],
    ['Resource access', 'Name, resource, member, and separate manage/effort flags.', 'Controls planning management and sensitive effort access.'],
], [1.38, 2.72, 2.73])
h('Why calendar versions matter')
p('Working calendars are fixed versions. Save a new version when the working schedule changes, then reference it from the applicable availability period. This preserves the meaning of earlier capacity records instead of rewriting the old schedule.')
h('Reading the date range')
p('Where a form says End date excluded, the final date is outside the range. For a period covering Monday through Friday, use Saturday as the excluded end. The timezone establishes the local working dates used by the calendar.')
lead('Example', 'A calendar can use Asia/Kolkata, eight hours Monday to Friday, zero hours on weekends, and the relevant holiday dates. A person\'s availability then applies that version over the selected period.')
lead('Long term value', 'Reliable calendars and availability make staffing estimates more realistic. They also preserve the assumptions behind decisions when schedules change.')

page('Resource commitments effort and capacity')
explain('Allocations',
    'An allocation reserves a resource\'s working time for a project.',
    'Name, resource, project, start date, excluded end date, and hours per scheduled day are required.',
    'The system needs the amount and period of each commitment to compare planned work with available time. Allocation management also requires the applicable resource and project permission.')
explain('Actual effort',
    'Actual effort records time already worked. It is separate from the allocation that was planned.',
    'Name, resource, project, work date, total hours, and included overtime hours are required.',
    'These entries support permitted actual utilization measures. Overtime is included within total hours, rather than added again. Sensitive effort requires a separate explicit permission, including for Admin.')
explain('Timesheet coverage',
    'Coverage indicates whether the resource\'s effort record is complete for a day.',
    'Name, resource, work date, and the Full resource timesheet is complete checkbox.',
    'A low recorded total may mean incomplete reporting. Actual utilization needs sufficient permissions and complete effort coverage across the selected period.')
explain('Allocation snapshots',
    'A snapshot preserves a dated view of the planned staffing.',
    'Name, resource, start date, and excluded end date.',
    'It creates a reference that can be reviewed later without relying only on the current editable allocations.')
h('What the capacity display shows')
p('Users select a resource and period, with supported skill, project, or team filters. The summary and heatmap show available hours, allocated hours, planned utilization, excess hours, and actual utilization where allowed. Permitted team workload signals summarize defined patterns; small groups and sensitive information are subject to privacy limits.')
lead('Example', 'Priya has eight available hours on Monday. Project A reserves six hours and Project B reserves four. The total allocation is ten hours, creating two excess hours and 125% planned utilization. The manager can move work or change staffing.')
p('Missing calendars, incomplete data, or insufficient scope can leave measures Unknown. Spare capacity on one day does not cancel an overload on a different day.')
lead('Long term value', 'Managers can identify repeated overbooking and plan staffing using actual commitments. This can reduce unrealistic schedules when the information is maintained across projects.')

page('7 Portfolios')
lead('Presentation sentence', 'This section groups projects so leaders can review the organization\'s overall investment and understand which business priorities the work supports.')
p('A project produces a specific result. A program coordinates related projects toward a shared result. A portfolio groups projects and programs for wider investment and priority discussions.')
table(['Record', 'Information to enter', 'Why it is needed'], [
    ['Portfolio', 'Required Name; description and owner.', 'Explains the group\'s purpose and accountability.'],
    ['Program', 'Required Name and Portfolio; description and owner.', 'Places related work within the wider investment group.'],
    ['Goal alignment', 'One project or portfolio scope; strategic goal; weight and contribution.', 'Records the claimed connection to organizational priorities. Weights for the selected scope must total 1, or 100%.'],
    ['Portfolio summary', 'Generated from permitted saved project records.', 'Shows supported financial and alignment information for the visible scope. Monetary totals remain separated by currency.'],
], [1.28, 2.75, 2.80])
h('Weight and contribution in simple words')
p('Weight means how much importance a goal receives in the alignment calculation. Contribution means how strongly the project or portfolio is recorded as supporting that goal. Both are entered judgments. The calculation does not independently prove business impact.')
h('A connected example')
p('A Customer Experience portfolio might contain a Customer Self Service program. The program could include a portal project and a support-process project. Their goal alignment can describe how they support the strategic goal of reducing customer effort.')
p('The hierarchy makes the reason for the work easier to explain: the portal is a delivery commitment, the program coordinates related changes, and the portfolio groups the investments for leadership review.')
h('How managers can use this over time')
bullets([
    'Keep the reason for each investment visible when staff or priorities change.',
    'Review competing projects together when money and people are limited.',
    'Discuss overlap between projects and check whether related efforts need coordination.',
    'Use recorded benefits and alignment to question whether continued investment remains justified.',
])
lead('Long term value', 'Portfolio records support more informed investment discussions. The quality of those discussions depends on accurate project, benefit, financial, and alignment information; the grouping alone does not choose the best investments automatically.')

page('8 Account and access')
lead('Presentation sentence', 'This section manages organization membership and personal accounts, while separate grants control who can work with each type of record.')
h('What is inside and what information is needed')
table(['Part', 'Information and function'], [
    ['Active organization', 'Select the organization you intend to work in. A person can have different roles in different organizations.'],
    ['Personal profile and security', 'Maintain display name and use the relevant password controls. Each security action requests the information it needs.'],
    ['Organization profile', 'Name, optional description, industry, and website describe the organization. Administrators and member counts are derived from active records.'],
    ['Invite a member', 'Admin enters the recipient\'s accessible email and intended role. An invitation joins an existing organization.'],
    ['Pending invitations', 'Shows recipient, role, expiry, and delivery state; Admin can refresh, resend, or cancel. A delivery request alone is not accepted membership.'],
    ['Organization members', 'Admin chooses role and Active status, then saves the member.'],
    ['Active sessions', 'Shows signed-in sessions and supports revoking relevant sessions.'],
    ['Access audit', 'Records membership and invitation actions for traceability.'],
], [1.8, 5.03])
h('Roles and responsibilities')
bullets([
    'Admin manages organization setup, membership, and grants. Sensitive effort still requires its separate permission.',
    'Project Manager works on permitted delivery records and can participate in assigned assessments.',
    'Viewer reads permitted delivery information and separately shared completed assessment scores. Viewer cannot submit answers, evidence, or reviews.',
])
h('Why a role is only part of access')
p('Project and portfolio grants, resource management, sensitive effort, framework editing, assessment assignments, and completed-result grants are separate controls. Project ownership and team membership do not automatically supply them. Read allows inspection; edit and manage provide the applicable permitted operations.')
lead('Example', 'A stakeholder can have Viewer membership, a project read grant, and a completed-result grant. They can inspect the project and shared scores without receiving private assessment evidence.')
p('Deactivating membership removes access to that organization. It does not delete the global account or its memberships elsewhere. The final active Admin is protected from removal without an appropriate handover.')
lead('Long term value', 'These controls support staff joining, changing duties, and leaving. Preserved access history helps explain who received permissions and when they changed.')

page('9 Component gallery and the connected workflow')
lead('Presentation sentence', 'The component gallery demonstrates reusable screen elements, helping the development team keep the application consistent as it grows.')
p('This section is mainly for development and design reference. Its sample records and numbers are illustrative examples, not live portfolio information.')
h('What is inside')
bullets([
    'Metric cards and badges demonstrate values and statuses.',
    'Tables demonstrate search, sorting, filters, and page navigation.',
    'Forms, drawers, pop-up windows, and confirmation messages demonstrate common interactions.',
    'An example chart includes its supporting table.',
    'Loading, empty, restricted-access, error, and stale-data examples show how common conditions should be explained.',
])
lead('Information to enter', 'The Example view name form accepts a name for the demonstration. Saving changes only that example view; it does not create a business record or permanently save organization data.')
lead('Long term value', 'Reusable patterns can reduce development and maintenance effort. Consistent screens also reduce how much users must relearn when a new module is added.')
h('How the sections work together')
p('Use a customer portal as one example throughout the presentation. The following sequence explains the relationship between the modules; organizations can adapt it to their own working process.')
numbered([
    'Account and access: invite the project manager, independent reviewer, and stakeholders.',
    'Organization setup: establish timezone, currency, teams, and strategic goals.',
    'Portfolios: place the portal within the Customer Experience portfolio.',
    'Projects: save the objective, tasks, milestones, money, risks, and reference plan.',
    'Resource planning: assign people using their schedules and other commitments.',
    'Integrations and jobs: import supported structured work records where needed.',
    'Maturity assessments: check whether the team follows consistent planning and review practices.',
    'Command center: review delivery, spending, risk, and staffing signals.',
    'Workspace overview: help each person find their responsibilities.',
    'Component gallery: help the development team keep future screens consistent.',
])

page('Presentation notes and simple terms')
h('Suggested opening')
p('Enterprise PPM connects our project plans, people, money, and management decisions. It gives us a shared place to record what we intend to deliver, monitor the work, and learn from the results. I will explain what each section does and why its information matters.')
h('Short speaking notes for each section')
table(['Section', 'Sentence to say'], [
    ['Workspace overview', 'This is our starting point and the place to find assigned responsibilities.'],
    ['Projects', 'Here we manage the individual project, including tasks, dates, costs, risks, and benefits.'],
    ['Maturity assessments', 'Here we check how consistently the team follows useful practices and review the evidence.'],
    ['Command center', 'Here managers review the overall position and investigate signals that need attention.'],
    ['Integrations and jobs', 'Here we bring supported work records into the system and check the import outcome.'],
    ['Resource planning', 'Here we compare people\'s available time with the work already assigned to them.'],
    ['Portfolios', 'Here we connect groups of projects to business priorities and review the wider investment.'],
    ['Account and access', 'Here we maintain membership and ensure permissions match responsibilities.'],
    ['Component gallery', 'Here the development team checks reusable patterns for consistent screens.'],
], [1.95, 4.88])
h('Terms the audience may ask about')
p('Governance means the rules for responsibility, decisions, and checks. A baseline is a fixed reference copy of a plan. A milestone is an important checkpoint. A sprint is a short planned work period. Mitigation is an action to reduce a risk. Allocation is time reserved for a project. Utilization compares assigned or actual work with available working time.')
h('Current capabilities and future plans')
p('The current foundation includes accounts and scoped access, delivery records, resource capacity, rubric assessments, defined dashboards, CSV imports, labelled integration fixtures, search, and in-app notices. Live Jira/GitHub acceptance still needs the relevant credentials and external checks. AI analysis, predictive models, broad document intake, and wider automated approvals, reminders, and report workflows belong to later planned work.')
h('Suggested closing')
p('Maintaining this information gives us a reliable record of what we planned, what happened, and where our practices need to improve. Over time, that record can support more realistic planning, clearer responsibility, and better informed decisions about future projects.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print('Words approximately', sum(len(x.text.split()) for x in doc.paragraphs) + sum(len(c.text.split()) for t in doc.tables for row in t.rows for c in row.cells))
