from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'learning' / 'Enterprise PPM Website User Testing Guide.docx'
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.7), Inches(8.3)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
    for border in st.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.07
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Enterprise PPM user testing guide  |  ').font.size = Pt(8)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)

def p(t, style=None):
    return doc.add_paragraph(t, style)
def h(t):
    doc.add_heading(t, 2)
def page(t):
    doc.add_page_break(); doc.add_heading(t, 1)
def steps(items):
    for i, item in enumerate(items, 1):
        a = doc.add_paragraph()
        a.add_run(f'{i}. ').bold = True
        a.add_run(item)
def check(t):
    a=p('Expected result  '); a.runs[0].bold=True; a.add_run(t)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,text in zip(t.rows[0].cells,headers):
        c.text=text
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EAEFF3'); c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs: r.bold=True
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=str(text)
    for row in t.rows:
        no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(3)
                a.paragraph_format.space_before=Pt(2)
                for r in a.runs: r.font.size=Pt(9.5)
    p('')
    return t

doc.add_heading('Enterprise PPM Website User Testing Guide',0)
p('Step by step checks with ready to enter sample data', 'Subtitle')
p('9 October 2026')
p('Use this guide to check the website as an ordinary user. Follow the pages in order: activate accounts, set up the organization, create projects, plan people, complete assessments and check dashboards. Each part tells you what to enter and what result to look for.')
h('Before you start')
steps([
    'Open /signup, enter your email and a fictional organization name, and choose Send verification instructions. Open the single-use link, set your display name and password, then sign in. In local mail mode, ask the operator to retrieve the private message; SMTP inbox delivery is still unverified. Existing accounts can use an Admin activation link.',
    'Open http://localhost:3000/login on the computer running the app. If you were given the separate demo website, use http://localhost:3002/login instead. Keep all testing on the same website.',
    'Use a separate test organization. The sample names below are fictional and are intended for practice. Add a new suffix such as Test 02 when repeating the guide.',
    'Use separate browser profiles for Admin, Project Manager and Viewer, or sign out before changing accounts. Two tabs in the same browser profile share the sign-in.',
    'Record Pass, Fail or Blocked for each check. Blocked means an account, permission or required record is missing; it does not automatically mean the feature is broken.'
])
h('Your route through the guide')
table(['Pages','What you will check'],[
    ['2 to 4','Accounts, invitations, organization, departments, teams and access'],
    ['5 to 7','Projects, tasks, milestones, spending, risks and benefits'],
    ['8 to 9','Working hours, staffing, time off and timesheets'],
    ['10 to 12','Maturity questions, answers, evidence, review and scores'],
    ['13 to 14','Dashboards, access restrictions and your test record'],
    ['15','One-time setup help for the person running the app']
],[1.1,5.6])
p('Sprint 8.1 also adds Getting started, My work, private notifications, organization name editing, record search, paged selectors, full-period comparisons and optional project and work metadata. Follow docs/SPRINT_8_1_USER_GUIDE.md for these checks and delivery configuration.')
p('The currently built screens cover organization access, delivery planning, resources, assessments and dashboards. AI chat, AI recommendations and automatic document scoring are planned features. Evidence upload currently accepts plain text files only.')

page('2 Activate accounts and sign in')
p('There is no preset password for the first Admin or the seeded demo account. The passwords in this table are examples to SET during activation; they will not work until the account exists and you have chosen that password. Use them only in the test organization.')
table(['Display name','Email address','Role','Example password to set'],[
    ['Asha Mehta','admin.ppm@example.com','Admin','PpmTest!Admin2026'],
    ['Rohan Shah','rohan.ppm@example.com','Project Manager','PpmTest!Rohan2026'],
    ['Meera Rao','meera.ppm@example.com','Project Manager','PpmTest!Meera2026'],
    ['Kiran Patel','viewer.ppm@example.com','Viewer','PpmTest!Viewer2026']
],[1.2,2.1,1.35,2.05])
h('Activate the first Admin')
steps([
    'Ask the person running the app to create Northstar Demo Services for admin.ppm@example.com and give you the latest activation link. Page 15 provides the setup instructions.',
    'Open that link. On Verify your account, enter Asha Mehta in Display name and PpmTest!Admin2026 in Password. The link normally fills the One-time token field.',
    'Click Verify and continue. Go to Sign in, enter admin.ppm@example.com and the password you just set, then click Sign in.',
    'Check the Account & access page. Confirm the email, Admin role and selected organization. Select Northstar Demo Services if you belong to more than one organization.'
])
check('The account opens successfully and the selected organization is correct. An activation link cannot be reused after it succeeds.')
h('If you are using the existing separate demo')
p('After the demo has been prepared, its first account is demo@example.com on port 3002. Ask for its activation link and choose a password there. There is no default demo password. other-demo@example.com belongs to the other demo organization. The sample accounts above are additional accounts to create, not existing demo logins.')
h('Check sign-in and recovery')
steps([
    'Sign out. Try the correct email with WrongPassword2026! once. Confirm sign-in fails, then sign in with the correct password.',
    'Open Account & access. Change Display name to Asha Mehta Test and click Save profile. Reload and confirm the new name remains. Change it back to Asha Mehta for the following steps.',
    'For a recovery check, use the Viewer account after creating it. Click Forgot password?, enter its email and click Send recovery instructions. Ask the app operator for the local recovery link.',
    'Open the recovery link within 30 minutes. Choose PpmTest!Viewer2027 and sign in with it. Confirm the old password stops working. Update your test notes with the new password.'
])
check('Wrong credentials are rejected, profile edits remain saved, and recovery replaces the password. New passwords need at least 12 characters; these examples meet that rule.')

page('3 Invite members and check their access')
p('Use Asha as Admin. Rohan will enter assessment answers; Meera will review them. They must be different active people. Kiran will check what a Viewer can see.')
h('Create and accept invitations')
steps([
    'Open Account & access. Under Invite a member, enter rohan.ppm@example.com, choose Project Manager and click Create invitation.',
    'Repeat for meera.ppm@example.com as Project Manager, then viewer.ppm@example.com as Viewer.',
    'Ask the person running the app for each invitation link from the local mailbox. This local build does not send these messages to real email inboxes.',
    'Open each link in the corresponding browser profile. Enter the display name and example password from page 2, then click Verify and continue.',
    'Sign in as each person once. As Asha, return to Organization members, reload if needed, and confirm all three memberships are active with the correct roles.'
])
check('All invited people can sign in. Rohan and Meera are Project Managers; Kiran is a Viewer. Only Admin sees membership administration controls.')
h('Sample access records to create after the projects exist')
p('A role alone does not give Rohan, Meera or Kiran access to every project. On Organization setup, use Project & portfolio access. Enter a Name for each record, select the Member, select the Project, leave Portfolio blank and choose Access. Do this after page 5, then reload each user’s website.')
table(['Name','Member','Project','Access'],[
    ['Rohan Portal Edit','Rohan Shah','Customer Portal Test','edit'],
    ['Rohan ERP Edit','Rohan Shah','ERP Upgrade Test','edit'],
    ['Meera Portal Read','Meera Rao','Customer Portal Test','read'],
    ['Kiran Portal Read','Kiran Patel','Customer Portal Test','read']
],[1.6,1.25,2.7,1.15])
p('Choosing a person as Owner or adding them to a team does not replace these access records. Assessment assignments and access to private timesheets are separate controls.')
h('Member changes')
steps([
    'As Admin, filter the member list for viewer.ppm@example.com. Open its membership edit control and confirm the role is Viewer.',
    'After the Viewer tests on page 14, deactivate Kiran and save. Reload Kiran’s open page and confirm organization access is refused.',
    'Reactivate Kiran and confirm access returns after a fresh sign-in if needed. Check the Admin organization history for the membership change.'
])
check('Membership changes take effect. Keep Asha active throughout this guide; the app protects the last active Admin from removal.')

page('4 Set up the organization and its structure')
p('Sign in as Asha. Use the tabs on Organization setup and Portfolios. For each record, click Create, fill Name and the fields below, then save. Leave unspecified optional fields blank. Choose records by name from the lists rather than typing internal record numbers.')
table(['Screen and tab','Name','Other sample fields'],[
    ['Organization setup\nOrganization configuration','Northstar Test Settings','IANA timezone: Asia/Kolkata\nCurrency: INR\nDescription: Website testing workspace'],
    ['Organization setup\nDepartments','Engineering Test','Owner: Asha Mehta\nDescription: Builds customer services'],
    ['Organization setup\nDepartments','Operations Test','Owner: Asha Mehta\nDescription: Runs internal services'],
    ['Organization setup\nTeams','Portal Team Test','Department: Engineering Test\nOwner: Rohan Shah'],
    ['Organization setup\nTeams','Operations Team Test','Department: Operations Test\nOwner: Meera Rao'],
    ['Organization setup\nStrategic goals','Improve Customer Service Test','Owner: Asha Mehta\nDescription: Reduce customer waiting time'],
    ['Portfolios\nPortfolios','Digital Transformation Test','Owner: Asha Mehta\nDescription: Customer and internal service improvements'],
    ['Portfolios\nPrograms','Customer Experience Test','Portfolio: Digital Transformation Test\nOwner: Asha Mehta']
],[1.7,2.1,2.9])
h('Add team membership')
steps([
    'Return to Organization setup and select Team membership.',
    'Create Rohan Portal Membership. Choose Portal Team Test and Rohan Shah. Set Start date to 2026-10-01 and leave End date blank.',
    'Create Meera Operations Membership. Choose Operations Team Test and Meera Rao, with the same start date.',
    'Reload the page. Check the department and team choices in the saved records.'
])
check('The departments, teams, portfolio and program remain after reload. Portal Team belongs to Engineering; Operations Team belongs to Operations. Team membership does not automatically give project edit access.')
p('Northstar Demo Services is the organization name created during initial setup. Organization configuration stores its timezone and currency; it is not a website screen for creating another organization.')

page('5 Create two projects')
p('Open Projects as Asha and click Create projects. Create both rows below. Use Digital Transformation Test as Portfolio and Customer Experience Test as Program for both projects.')
table(['Field','Customer Portal Test','ERP Upgrade Test'],[
    ['Description','Self-service customer website','Upgrade internal finance tools'],
    ['Department','Engineering Test','Operations Test'],
    ['Delivery team','Portal Team Test','Operations Team Test'],
    ['Owner','Rohan Shah','Meera Rao'],
    ['Status','active','active'],
    ['Reported health','Green','Red'],
    ['Currency','INR','INR'],
    ['Start date','2026-10-01','2026-10-01'],
    ['Planned finish','2026-12-01','2026-12-01'],
    ['Forecast finish','2026-11-25','2027-01-15'],
    ['Actual start','2026-10-01','2026-10-01'],
    ['Actual finish','Leave blank','Leave blank'],
    ['Expenditure ledger is complete','Leave unchecked for now','Leave unchecked']
],[1.9,2.4,2.4])
h('Check the project list and saved information')
steps([
    'Save each project and reload Projects. Search records for Customer Portal Test. Confirm the matching project appears.',
    'Change Filter status to active and check the result. Clear the search, try Sort records, and confirm both projects remain available.',
    'Open Customer Portal Test from the list. Confirm its reported health, dates, team and currency match the table.',
    'Return to Projects, edit its Description to Self-service customer website with online support and save. Open History for that record and check the change.',
    'Create the project access records from page 3. Sign in as Rohan, reload and confirm both projects are visible. Kiran should see Customer Portal Test only.'
])
check('Both projects save correctly. Search and filters show the matching records. The forecast for ERP is later than its plan. Each person sees the projects covered by their access.')
h('How to use the next pages')
p('Most delivery records are created inside a project. Open Customer Portal Test, choose the relevant tab, then click its Create button. Every record has a Name; use the names given below so you can find them again. Use Asha for the full walkthrough, then repeat a small edit as Rohan to check his permission.')

page('6 Add tasks milestones and delivery updates')
p('Open Customer Portal Test. Create the sprint before adding its work items. A sprint is a short period of planned work; points are the team’s estimate of task size.')
table(['Tab and name','Sample values'],[
    ['Sprints\nPortal Sprint 1 Test','Team: Portal Team Test\nStart: 2026-10-12; End: 2026-10-23; Status: planned'],
    ['Work items\nBuild Sign In Test','Sprint: Portal Sprint 1 Test; Owner: Rohan Shah\nStatus: ready; Points: 5; Start: 2026-10-12; Due: 2026-10-16'],
    ['Work items\nBuild Profile Page Test','Sprint: Portal Sprint 1 Test; Owner: Rohan Shah\nStatus: in_progress; Points: 3; Due: 2026-10-20'],
    ['Work items\nCheck Help Page Test','Sprint: Portal Sprint 1 Test\nStatus: done; Points: 2; Due: 2026-10-15'],
    ['Milestones\nPortal Design Approved Test','Due: 2026-10-16; Actual finish: 2026-10-15\nCritical milestone: unchecked'],
    ['Milestones\nPortal Launch Test','Due: 2026-12-01; Actual finish: blank\nCritical milestone: checked'],
    ['Blockers\nWaiting For Test Access','Related work item: Build Sign In Test; Owner: Rohan Shah\nStatus: open; Description: Test account access is pending'],
    ['Status updates\nPortal Weekly Update Test','Reported health: Amber\nDescription: Sign-in testing is waiting for test account access']
],[2.3,4.4])
h('Check updates and saved history')
steps([
    'On Work items, edit Build Sign In Test to in_progress and tick Blocked. Save. Check that it appears in the correct Work board column.',
    'Open its History and confirm the previous and new values are visible. Check the Delivery timeline for task and milestone dates.',
    'Resolve Waiting For Test Access. Edit Build Sign In Test, untick Blocked and set Status to done. Confirm the board changes after saving.',
    'On Plan baselines, create Portal Original Plan Test. Enter Baseline reason: Initial dates and scope agreed for website testing. Save and view the saved plan.',
    'Edit Portal Launch Test due date to 2026-12-04. Check that the current timeline changes and the saved baseline still shows the original date.',
    'Close Portal Sprint 1 Test only after recording the tasks you want completed. First edit a planned sprint to active. Close it only after its required work is completed. A closed sprint is immutable; do not attempt to reopen it. Reopening completed work instead requires a Change reason.'
])
check('Task status changes are reflected on the board. Changes have history. A saved baseline keeps the earlier plan even when the live dates change.')
p('Optional dependency check: on Dependencies, create Portal Before ERP Test with Customer Portal Test as Project and ERP Upgrade Test as Successor project. Leave task and milestone links blank for a project-level link. A reverse ERP-to-Portal link should be rejected because it would create a circular dependency.')

page('7 Add budgets risks and business benefits')
p('Open Customer Portal Test as Asha. These amounts all use INR so they can be compared directly. Use the corresponding tabs to create the records.')
table(['Tab and name','Sample values'],[
    ['Budget lines\nPortal Build Budget Test','Amount: 500000; Currency: INR; Approved budget: checked'],
    ['Budget lines\nPortal Support Budget Test','Amount: 100000; Currency: INR; Approved budget: checked'],
    ['Expenditure\nPortal Design Payment Test','Amount: 80000; Currency: INR; Expenditure date: 2026-10-07\nReverse original expenditure: blank'],
    ['Expenditure\nPortal Build Payment Test','Amount: 120000; Currency: INR; Expenditure date: 2026-10-08\nReverse original expenditure: blank'],
    ['Risk register\nSupplier Delay Test','Owner: Rohan Shah; Severity: 4; Likelihood: 3\nStatus: open; Review date: 2026-10-20\nMitigation: Confirm delivery weekly and keep a backup supplier'],
    ['Business benefits\nCustomer Waiting Time Test','Unit: hours saved; Target: 100; Realized: 25\nPeriod start: 2026-10-01; Period end excluded: 2026-11-01'],
    ['Goal alignment\nPortal Customer Goal Test','Strategic goal: Improve Customer Service Test\nWeight: 1; Contribution: 0.8'],
    ['Metric plans\nPortal Forecast Test','Frozen baseline: Portal Original Plan Test\nWorking calendar: India Work Week Test (create on page 8 first)\nForecast final cost: 550000']
],[2.3,4.4])
h('Check the financial information')
steps([
    'After entering both expenditure records, edit the Customer Portal Test project and check Expenditure ledger is complete. This means the spending records are complete for this sample.',
    'Reopen the project overview. Approved budget should be 600000 INR and Recorded expenditure should be 200000 INR. The simple remaining amount is 400000 INR.',
    'Check the history of a budget edit. Expenditure is a permanent entry; do not expect an ordinary Edit button for it.',
    'Optional correction test: create Portal Payment Reversal Test with Amount -120000, Currency INR, date 2026-10-09 and Reverse original expenditure set to Portal Build Payment Test. Spending should then be 80000 INR. Skip this if you want to keep the main 200000 result.',
    'Add a budget of 900000 INR and one expenditure of 300000 INR to ERP Upgrade Test. Keep its ledger incomplete to check how missing information appears on dashboards.'
])
check('Totals match the saved amounts. A reversal keeps the original entry visible. ERP spending-based measures may remain unknown or partial because its ledger is incomplete.')
p('A realized benefit of 25 against a target of 100 represents 25% achievement. Closing a risk requires a resolution or acceptance reason; use Backup supplier confirmed. Complete the Metric plans record after creating its calendar on the next page.')

page('8 Plan people and their available hours')
p('Open Resource planning as Asha. A resource profile represents a person whose time you plan; it can exist without a sign-in account. Use Planning records to choose the record type, then click Create.')
table(['Planning records and name','Sample values'],[
    ['Working calendars\nIndia Work Week Test','IANA timezone: Asia/Kolkata\nMonday to Friday: 8 hours each; Saturday and Sunday: 0\nHoliday dates: leave blank'],
    ['Resource profiles\nRohan Planner Test','Login member: Rohan Shah'],
    ['Resource profiles\nPriya Designer Test','Login member: leave blank'],
    ['Availability\nRohan October Availability Test','Resource: Rohan Planner Test; Calendar version: India Work Week Test\nStart: 2026-10-12; End excluded: 2026-10-17'],
    ['Resource teams\nRohan Portal Resource Team Test','Resource: Rohan Planner Test; Team: Portal Team Test\nStart: 2026-10-12; End excluded: 2026-10-17'],
    ['Skills\nRohan Testing Skill Test','Resource: Rohan Planner Test; Skill: Website Testing'],
    ['Allocations\nRohan Portal Allocation Test','Resource: Rohan Planner Test; Project: Customer Portal Test\nStart: 2026-10-12; End excluded: 2026-10-17; Hours per day: 6'],
    ['Allocations\nRohan ERP Allocation Test','Resource: Rohan Planner Test; Project: ERP Upgrade Test\nSame dates; Hours per day: 4']
],[2.25,4.45])
h('Check the workload numbers')
steps([
    'Choose Rohan Planner Test as Resource profile. Set From date to 2026-10-12 and Until date excluded to 2026-10-17. Use All authorized projects for the project slice.',
    'Click Refresh staffing if needed. Check Available hours: 40, Allocated hours: 50, Planned utilization: 125% and Excess hours: 10.',
    'Check each of the five working days. Each should show 8 available hours and 10 allocated hours, so the person is over capacity.',
    'Edit Rohan ERP Allocation Test to 2 hours per day. Refresh and confirm 40 allocated hours, 100% planned utilization and no excess.',
    'Create Approved time off named Rohan Half Day Test for Rohan, start 2026-10-14, end excluded 2026-10-15 and 4 hours per scheduled day. Available hours should become 36; allocation stays 40, leaving 4 excess hours.'
])
check('The summary, daily table and workload colors agree. An end date marked excluded is not counted: the first interval covers 12 to 16 October, not 17 October.')
p('Leave the ERP allocation at 2 hours and the half day off in place for page 9. Return to page 7 and create Portal Forecast Test now that India Work Week Test exists.')

page('9 Enter actual hours and preserve a plan')
p('Timesheets are private. Even an Admin needs explicit sensitive effort access to read or enter these records. Keep the five-day period from page 8 selected.')
h('Give the test users the required resource access')
steps([
    'As Asha, select Resource access under Planning records. Create Asha Rohan Effort Test with Resource Rohan Planner Test, Member Asha Mehta and May access sensitive effort checked.',
    'Create Rohan Own Planning Test for Rohan Planner Test and Rohan Shah. Check May manage resource and May access sensitive effort.',
    'Refresh staffing. Keep Rohan Planner Test selected. Rohan also needs edit access to each project whose allocation he changes; page 3 gives him both grants.'
])
h('Create actual effort and mark each day complete')
p('Choose Actual effort. Each row needs a Name, Resource, Project, Worked on, Total hours and Included overtime hours. Use Rohan Planner Test for every row. Enter overtime as 0.')
table(['Worked on','Portal hours','ERP hours','Suggested name endings'],[
    ['2026-10-12','6','2','Oct12 Portal / Oct12 ERP'],
    ['2026-10-13','6','2','Oct13 Portal / Oct13 ERP'],
    ['2026-10-14','4','0','Oct14 Portal only'],
    ['2026-10-15','6','2','Oct15 Portal / Oct15 ERP'],
    ['2026-10-16','6','2','Oct16 Portal / Oct16 ERP']
],[1.3,1.1,1.1,3.2])
steps([
    'Create the positive-hour entries above using Customer Portal Test or ERP Upgrade Test as appropriate. Use names such as Rohan Oct12 Portal Test. Do not create the zero-hour ERP entry.',
    'Select Timesheet coverage. Create one record for each of the five dates, named Rohan Oct12 Coverage Test and so on. Choose Rohan Planner Test and tick Full resource timesheet is complete.',
    'Refresh staffing. The total actual hours are 36. With 36 available hours, Actual utilization should be 100% for the full resource period.',
    'Untick completeness for one day and save. Confirm actual utilization becomes unknown or partial. Restore completeness and check it again.'
])
check('The app distinguishes planned hours from hours actually worked. Missing completeness is not silently treated as zero work.')
h('Check an allocation snapshot')
steps([
    'Under Allocation snapshots, create Rohan Week Plan Test for Rohan Planner Test, start 2026-10-12 and end excluded 2026-10-17.',
    'Click View snapshot. Change a current allocation, then view the snapshot again. Confirm its saved planning inputs retain their earlier values. Restore the allocation to 2 ERP hours per day when finished.'
])

page('10 Prepare maturity assessments')
p('A framework is a set of questions. A published version fixes that question set for an assessment. The respondent answers it; the independent reviewer checks and approves it. Use the built-in PMO example for a complete first test.')
h('Load publish and assign')
steps([
    'As Asha, open Maturity assessments. Click Add example frameworks. Confirm four examples appear in Framework library.',
    'Find PMO Capability (Example). Click Edit draft and inspect the questions. Keep their values unchanged and close the editor. Click Publish version once and note the version number.',
    'Click Assign assessment on that framework. Enter Engineering PMO October Test. Select the published version, Department Engineering Test, Team Portal Team Test, Respondent Rohan Shah and Independent reviewer Meera Rao. Click Create assignment.',
    'Assign a second assessment named Operations PMO October Test using exactly the same published version. Select Operations Test and Operations Team Test, respondent Rohan and reviewer Meera.',
    'Sign in as Rohan and open Maturity assessments. Open Engineering PMO October Test. Answer every question using the Engineering column below.'
])
h('PMO question and answer matrix')
table(['Area and share','Question shown on website','Engineering answer','Operations answer'],[
    ['Governance\n40% overall','Project charter\nIs a project charter required and reviewed for all projects?','2 Defined','1 Informal'],
    ['Governance\n40% overall','Portfolio alignment\nAre projects formally linked to strategic goals with documented weights?','Yes','No'],
    ['Delivery\n60% overall','Baseline tracking\nHow consistently are project baselines tracked and updated?','3 All','1 Occasional']
],[1.2,3.1,1.2,1.2])
p('Project charter choices are 0 None, 1 Informal, 2 Defined and 3 Enforced. Baseline tracking choices are 0 None, 1 Occasional, 2 Most and 3 All. All three questions have equal question weight within their own area.')
h('Evidence notes to enter')
table(['Question','Engineering evidence note'],[
    ['Project charter','The portal charter uses the standard template and was reviewed on 5 October 2026.'],
    ['Portfolio alignment','The portal project is linked to Improve Customer Service Test with weight 1.'],
    ['Baseline tracking','Portal Original Plan Test was saved and compared with updated milestone dates.']
],[1.65,5.05])
check('Engineering should show Governance 83.33, Delivery 100.00 and overall 93.33 out of 100 before reviewer changes. Operations should show Governance 16.67, Delivery 33.33 and overall 26.67. Scores remain provisional until review is completed.')

page('11 Save submit and review an assessment')
h('As Rohan save answers and add a text file')
steps([
    'Enter the Engineering answers and evidence notes from page 10. Let any automatic save finish, or click Save answers when available. Reload and reopen the assessment to confirm the answers remain.',
    'Create a plain text file named portal-charter-test.txt in Notepad. Paste the sample text below and save it as UTF-8. Attach it under Project charter after answers are saved.',
    'Confirm the file is listed as quarantined. That means it is waiting for the reviewer to check and release it. It should not be offered for download yet.',
    'Enter Review or submission note: Engineering answers and sample evidence are ready for review. Click Submit for review. If it is disabled, wait for the answer save to finish.',
    'Confirm the status changes to submitted. Respondent answer editing should now be unavailable.'
])
p('Sample text for portal-charter-test.txt:')
p('Project: Customer Portal Test\nDocument: Sample project charter\nPurpose: Let customers use online support services\nOwner: Rohan Shah\nReview date: 2026-10-05\nStatus: Standard template reviewed\nEvidence: Training sample for website testing')
h('As Meera review and complete')
steps([
    'Sign in as Meera in her browser profile. Open Engineering PMO October Test. Enter Review or submission note: Reviewing answers against the sample records. Click Start review.',
    'Click Preview portal-charter-test.txt and read the plain text. Enter a release note such as Reviewed the training text and confirmed it contains the stated charter details. Click Release portal-charter-test.txt.',
    'Confirm the released file can now be downloaded by an authorized user. Try Complete review only after every attached file has been reviewed and released.',
    'Optional score adjustment: enter 33.3333 in Adjusted score for Project charter and enter reason Standard template exists but is not yet used consistently. This field uses a score out of 100, rather than the question level. The adjustment is saved when you complete review.',
    'Enter the final note: Reviewed the sample answers and released the evidence. Click Complete review. Reopen the assessment and check the final result and History.'
])
check('The approved assessment is completed and cannot be edited. History records submission, review, evidence release and any score adjustment. With the optional adjustment, Governance is 66.67 and overall is 86.67. Without it, the Engineering result remains 93.33.')
h('Check return and resubmission using Operations')
p('Rohan enters the Operations answers from page 10 and submits. Meera adds Please add evidence notes for each answer and clicks Return for changes. Rohan adds the notes, saves and submits again. Meera completes review. Department comparison should now contain both completed assessments because they use the same published version.')
p('Upload limits: plain text .txt, nonempty, up to 512 KiB per file and 50 files per assessment. PDF, Word, images and recordings are not supported by this intake. A reviewer must release attached files before completion.')

page('12 Test the other question sets')
p('These are the other built-in product examples. They are practice models, not certified industry assessments. Publish each once, assign it to Rohan with Meera as reviewer, enter the sample answers and complete the review without adjustments to check the scores below.')
examples=json.loads((ROOT/'services/api/ppm_api/assessment_examples.json').read_text())
answers={'Sprint ceremonies':'2 Consistent','Velocity stability':'Yes','Policy review cadence':'2 Annual','Estimation accuracy':'3 Calibrated','Milestone adherence':'2 70-90%'}
for ex in examples[1:]:
    h(ex['name'].replace('&','and').replace('(','').replace(')',''))
    rows=[]
    for d in ex['dimensions']:
        for q in d['questions']:
            rows.append([d['name']+'\n'+str(int(float(d['weight'])*100))+'% overall',q['text'],answers[q['name']]])
    table(['Area and share','Question','Sample answer'],rows,[1.5,3.9,1.3])
    if ex['category']=='agile':
        p('Choices for Sprint ceremonies: 0 None, 1 Partial, 2 Consistent, 3 Improved. Velocity stability is Yes or No. Expected overall score: 83.33. Sample evidence note: Planning, reviews and retrospectives were held; four sprint totals were recorded.')
    elif ex['category']=='governance':
        p('Choices: 0 Never, 1 Ad hoc, 2 Annual, 3 Continuous. Expected overall score: 66.67. Sample evidence note: The governance policy was reviewed on 1 October 2026 and the next annual review is scheduled.')
    else:
        p('Estimation choices: 0 None, 1 Occasional, 2 Regular, 3 Calibrated. Milestone choices: 0 <50%, 1 50-70%, 2 70-90%, 3 >90%. Expected overall score: 83.33. Sample evidence note: Estimates are compared with actual results; 8 out of 10 milestones met the original dates.')
h('Optional custom question set')
p('As Asha, click Create framework. Name it Simple Readiness Test, category custom. Keep one area named Readiness with weight 1. Add a Yes / no question named Plan approved with prompt Is the project plan approved?, weight 1 and Required answer checked. Add a second question named Next action, type Text (not scored), weight 1 and Required answer checked. Prompt: What will the team do next? Save framework draft and Publish version.')
p('Assign it to Rohan with Meera as reviewer. Answer Yes and Confirm test access on 20 October 2026. Expected score: 100. The text answer is required but does not increase the score. Leaving it blank should prevent submission. This also checks that question sets can be created through the website.')

page('13 Check the command center dashboards')
p('Use Asha first. Complete the sample records before expecting useful dashboard values. A dashboard with an unknown value may be correctly showing missing information.')
steps([
    'Open Command center. Set From date to 2026-10-01 and To date excluded to 2026-11-01. Choose Digital Transformation Test as Portfolio and click Refresh.',
    'Open Executive, Delivery, Risk and Resource in turn. Confirm each opens and displays only projects you are allowed to see.',
    'For a displayed measure, click Inspect followed by its label. Check the source information, period and explanation. Confirm the numbers can be traced to your saved records.',
    'Enter Northstar October Test View in New view name and click Save view. Change the filters, then choose that name from Personal saved view. Confirm the saved dates, portfolio and dashboard return.',
    'Edit one of your source records, return here and click Refresh. Confirm the affected value updates. Saved views store filters; they do not preserve a frozen copy of the numbers.'
])
table(['Dashboard','What to check with these records'],[
    ['Executive','Portal approved budget 600000 INR and recorded spending 200000 INR if you skipped the reversal. Do not confuse spending with forecast cost. Benefits show the saved target and realized values.'],
    ['Delivery','Portal tasks, launch milestone, blocker and forecast dates agree with project records. After sprint closure, inspect the closed sprint result. Team points are estimates, not hours.'],
    ['Risk','Supplier Delay Test appears with severity 4, likelihood 3, owner, mitigation and review date. Check a saved status change after Refresh. Reported project health can differ from calculated risk.'],
    ['Resource','Rohan’s visible planning data appears. For the exact 36 available and 40 allocated hour check, use 2026-10-12 to 2026-10-17 excluded instead of the full month. A full month can be unknown because availability was entered only for five days.']
],[1.2,5.5])
h('Check incomplete information')
p('ERP Upgrade Test deliberately has incomplete spending records. Confirm the website does not present unknown spending-based performance as a reliable zero. A measure can also be unknown if the baseline, working calendar, forecast, completed sprint or full timesheet coverage is missing. Read its explanation before recording a failure.')
h('Check access and saved views')
p('Open Command center as Kiran. Only Customer Portal Test should be visible after his read grant. ERP names and details should not appear. Asha’s personal saved view should not be offered to Kiran. Saving a view does not give a user new project access.')
check('Filters, the four dashboards and Inspect controls work. Source edits appear after refresh. Missing information is explained, and each user sees only authorized data.')

page('14 Final checks and a test record')
h('Check restrictions and everyday use')
table(['Check','Action','Expected result'],[
    ['Viewer','As Kiran, open Customer Portal Test and try to edit a task, budget or risk.','Write controls are unavailable or the change is refused. Read access still works.'],
    ['Separate review','As Rohan, open his submitted assessment.','He cannot approve his own submission; Meera performs the review.'],
    ['Required answers','Use a new assessment. Leave a required question blank and try to submit.','Submission is refused with a useful message.'],
    ['Published questions','Edit a framework draft after assigning an assessment.','The existing assessment keeps its original question set.'],
    ['Private hours','Use Kiran or Meera without a sensitive effort grant.','Private timesheet entries and details are unavailable.'],
    ['Old page','Edit the same task in two tabs; save one, then try saving the older form.','A stale change is refused. Reload and confirm the latest values before editing again.'],
    ['Small screen','Narrow the browser window. Open the navigation menu, project and assessment.','Controls remain usable; wide tables can be scrolled without hiding their content.'],
    ['Navigation','Use workspace command search, open a destination and switch theme.','Navigation opens the chosen screen and the theme changes.'],
    ['Sign out','Use Account & access to sign out, then reopen a project link.','The website requires sign-in before showing protected records.']
],[1.2,2.8,2.7])
p('If you have access to two test organizations, switch between them using Organization. Records from one should not appear in the other. With only one organization available, mark this check Blocked rather than guessing.')
h('If a step cannot continue')
p('No invitation or recovery email: ask the app operator for the private local mailbox link. No project or edit button: confirm the active organization, role and project access, then reload. Submit disabled: let answer saving finish. Complete review refused: check required answers and release every attachment. Unknown capacity: check the selected dates and availability records.')
h('Copy this record for each issue')
table(['Field','Your test notes'],[
    ['Date and tester','____________________________________________'],
    ['Website and organization','____________________________________________'],
    ['Account role and check','____________________________________________'],
    ['Steps and sample record name','____________________________________________'],
    ['Expected and actual result','____________________________________________'],
    ['Result','Pass / Fail / Blocked'],
    ['Issue reference and follow-up','____________________________________________']
],[2.1,4.6])
p('Use sample data in screenshots. Do not include passwords, activation links or recovery links in issue reports. Mark the guide complete when accounts, organization, project work, budgets, staffing, assessments and dashboards have each been checked.')

page('15 One time setup help for the app operator')
p('This page is for the person running the app. The tester can follow pages 2 to 14 in the browser after receiving working links. Run these commands from the enterprise-ppm project folder on the normal local installation, with Docker running and the app already started.')
h('Create the test organization and first Admin')
p('For a new identity setup:')
p('docker compose run --rm migrate python scripts/bootstrap_identity.py --email admin.ppm@example.com --organization "Northstar Demo Services"')
p('If an organization already exists and you want a separate test organization, use the explicit additional-organization option:')
p('docker compose run --rm migrate python scripts/bootstrap_identity.py --email admin.ppm@example.com --organization "Northstar Demo Services" --additional-organization')
p('Retrieve the activation message privately and give the latest unexpired activation link to the tester:')
p('docker compose exec -T api python scripts/local_mailbox.py --email admin.ppm@example.com --kind activate')
h('Retrieve invitations after Asha creates them')
for email in ['rohan.ppm@example.com','meera.ppm@example.com','viewer.ppm@example.com']:
    p(f'docker compose exec -T api python scripts/local_mailbox.py --email {email} --kind invite')
p('Activation and invitation links expire after 24 hours. Use Resend invitation on Account and access if needed; the previous link becomes invalid. Cancel invitation invalidates every outstanding link. Use the existing identity setup and recovery instructions if an activation link has expired; do not assume rerunning a seed issues a new link.')
h('Retrieve a password recovery message')
p('After the user clicks Send recovery instructions:')
p('docker compose exec -T api python scripts/local_mailbox.py --email viewer.ppm@example.com --kind reset')
p('Recovery links expire after 30 minutes and work once. These messages stay in the local app mailbox; no real email delivery is configured for this workflow.')
h('If using the isolated demo on port 3002')
p('Use the compose.demo.yaml commands and matching demo configuration described in docs/DEMO-DEVELOPMENT.md. The demo database and accounts are separate from the normal port 3000 installation. Retrieve its message with:')
p('docker compose -f compose.yaml -f compose.demo.yaml exec -T api python scripts/local_mailbox.py --email demo@example.com --kind activate')
p('Keep private challenge links out of shared documentation. None of the example passwords in this guide are existing account credentials; the user sets a password during activation.')
h('Project references for maintaining this guide')
p('Screen names and sample fields follow the current website components: identity-client.tsx, delivery.tsx, delivery-fields.ts, capacity.tsx, assessments.tsx and command-center.tsx. The built-in questions come from services/api/ppm_api/assessment_examples.json. Account setup and review rules follow docs/IDENTITY-DEVELOPMENT.md and docs/contracts/assessment.md.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.title='Enterprise PPM Website User Testing Guide'
doc.core_properties.subject='Step by step website checks and fictional sample data'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)
